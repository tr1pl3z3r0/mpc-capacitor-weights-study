"""
fase5_excel.py — Escribe el Excel final (sección 7, Fase 5), a partir de
resultados/fase4_seleccion.json y resultados/simulaciones.csv.

Trabaja sobre resultados/Pruebas_por_peso_MPC_resultados.xlsx (copia de
trabajo, NUNCA la plantilla original). Conserva formato y estructura.

Mapa de bloques verificado contra la plantilla real (sección 7, Fase 5):
  bloque k (k=0..4, uno por horizonte): empieza en fila r = 3 + 15*k
    r      : título (añadir "Horizonte N=<h> con los pesos...")
    r+3..9 : q11, q22, q33, q44, q55, r11, r22
    r+11..14: Osc. Capacitores (V), RMS corriente circ. (A),
              Max corriente circ. (A), Min corriente circ. (A)
  columnas C..L = Pruebas 1..10

Uso:
    python fase5_excel.py
"""

import json

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import config

VERDE = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ROJO = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FUENTE_NEGRITA = Font(bold=True)

FILAS_PESOS = ["q11", "q22", "q33", "q44", "q55", "r11", "r22"]
VCAP_NOMBRES = ["ap", "bp", "cp", "an", "bn", "cn"]


def _to_float(row: dict, key: str, default=None):
    v = row.get(key, "")
    if v in ("", None):
        return default
    try:
        return float(v)
    except (ValueError, TypeError):
        return default


def _col_letter(prueba_num: int) -> str:
    """Prueba 1 -> C, Prueba 2 -> D, ..., Prueba 10 -> L."""
    return get_column_letter(2 + prueba_num)


def escribir_bloque(ws, k: int, horizonte: int, bloque: dict):
    r = 3 + 15 * k
    pruebas = [bloque["prueba_1_caso_base"]] + bloque["pruebas_2_a_10"]
    no_cumple_ids = set(bloque.get("no_cumple_ids", []))

    # Verificar etiquetas antes de escribir (sección 7: "comprueba que coinciden").
    etiquetas_esperadas = {
        r: "Horizonte",  # prefijo, el texto completo varía
        r + 1: "Prueba",
        r + 3: "q11", r + 4: "q22", r + 5: "q33", r + 6: "q44", r + 7: "q55",
        r + 8: "r11", r + 9: "r22",
        r + 11: "Osc. Capacitores (V)",
        r + 12: "RMS corriente circ. (A)",
        r + 13: "Max corriente circ. (A)",
        r + 14: "Min corriente circ. (A)",
    }
    for fila_check, prefijo in etiquetas_esperadas.items():
        valor = ws.cell(row=fila_check, column=2).value or ""
        if not str(valor).strip().startswith(prefijo):
            raise RuntimeError(
                f"Fila {fila_check} col B esperaba prefijo {prefijo!r}, encontró {valor!r}. "
                f"El mapa de filas no coincide con la plantilla."
            )

    ws.cell(row=r, column=2).value = (
        f"Horizonte N = {horizonte} con los pesos del control de las corrientes "
        f"en Q_Val = {config.Q_VAL}, R_Val = {config.R_VAL}"
    )

    mejor_rms = None
    mejor_idx = None
    for i, p in enumerate(pruebas):
        if p is bloque["prueba_1_caso_base"]:
            continue
        if p["id"] in no_cumple_ids:
            continue
        rms = _to_float(p, "rms_icirc_peor")
        if rms is not None and (mejor_rms is None or rms < mejor_rms):
            mejor_rms = rms
            mejor_idx = i

    for i, prueba in enumerate(pruebas):
        prueba_num = i + 1
        col = _col_letter(prueba_num)

        for j, peso in enumerate(FILAS_PESOS):
            fila_peso = r + 3 + j
            celda = ws[f"{col}{fila_peso}"]
            celda.value = _to_float(prueba, peso)
            celda.number_format = "0.00E+00"

        fila_osc = r + 11
        fila_rms = r + 12
        fila_max = r + 13
        fila_min = r + 14

        osc_max = _to_float(prueba, "osc_max")
        rms_peor = _to_float(prueba, "rms_icirc_peor")
        max_icirc = _to_float(prueba, "max_icirc")
        min_icirc = _to_float(prueba, "min_icirc")

        c_osc = ws[f"{col}{fila_osc}"]
        c_osc.value = osc_max
        c_osc.number_format = "0.000"
        if osc_max is not None:
            c_osc.fill = VERDE if osc_max < config.OSC_MAX else ROJO

        for fila_m, valor in [(fila_rms, rms_peor), (fila_max, max_icirc), (fila_min, min_icirc)]:
            c = ws[f"{col}{fila_m}"]
            c.value = valor
            c.number_format = "0.000"

        if i == mejor_idx:
            for fila_m in [fila_osc, fila_rms, fila_max, fila_min] + [r + 3 + j for j in range(7)]:
                ws[f"{col}{fila_m}"].font = FUENTE_NEGRITA


def escribir_hoja_metricas_adicionales(wb, seleccion: dict):
    nombre = "Métricas adicionales"
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)

    headers = (
        ["horizonte", "prueba_num", "sim_id", "estado"]
        + [f"osc_vcap_{n}" for n in VCAP_NOMBRES]
        + ["desv_max"]
        + [f"media_vcap_{n}" for n in VCAP_NOMBRES]
        + ["rmse_vcap", "itae", "pico_abs_icirc", "idc_abs_max", "margen_10v", "no_cumple"]
    )
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h).font = FUENTE_NEGRITA

    fila_out = 2
    for h_str, bloque in seleccion.items():
        no_cumple_ids = set(bloque.get("no_cumple_ids", []))
        pruebas = [bloque["prueba_1_caso_base"]] + bloque["pruebas_2_a_10"]
        for i, p in enumerate(pruebas, start=1):
            osc_max = _to_float(p, "osc_max")
            margen = (config.OSC_MAX - osc_max) if osc_max is not None else None
            fila_vals = (
                [int(h_str), i, p.get("id"), p.get("estado")]
                + [_to_float(p, f"osc_vcap_{n}") for n in VCAP_NOMBRES]
                + [_to_float(p, "desv_max")]
                + [_to_float(p, f"media_vcap_{n}") for n in VCAP_NOMBRES]
                + [_to_float(p, "rmse_vcap"), _to_float(p, "itae"), _to_float(p, "pico_abs_icirc"),
                   _to_float(p, "idc_abs_max"), margen, p.get("id") in no_cumple_ids]
            )
            for j, v in enumerate(fila_vals, start=1):
                ws.cell(row=fila_out, column=j, value=v)
            fila_out += 1


def escribir_hoja_resumen(wb, seleccion: dict):
    nombre = "Resumen"
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)

    headers = ["horizonte", "n_simulaciones", "n_factibles", "mejor_sim_id",
               "mejor_rms_icirc_peor", "mejor_osc_max"] + FILAS_PESOS
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h).font = FUENTE_NEGRITA

    fila_out = 2
    for h_str, bloque in seleccion.items():
        no_cumple_ids = set(bloque.get("no_cumple_ids", []))
        candidatos = [p for p in bloque["pruebas_2_a_10"] if p["id"] not in no_cumple_ids]
        mejor = candidatos[0] if candidatos else None

        fila_vals = [int(h_str), bloque["n_simulaciones"], bloque["n_factibles"]]
        if mejor:
            fila_vals += [mejor.get("id"), _to_float(mejor, "rms_icirc_peor"), _to_float(mejor, "osc_max")]
            fila_vals += [_to_float(mejor, p) for p in FILAS_PESOS]
        else:
            fila_vals += [None, None, None] + [None] * len(FILAS_PESOS)

        for j, v in enumerate(fila_vals, start=1):
            ws.cell(row=fila_out, column=j, value=v)
        fila_out += 1


def escribir_hoja_sensibilidad(wb):
    nombre = "Sensibilidad"
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)

    conclusiones_path = config.RESULTADOS_DIR / "fase2_conclusiones.json"
    headers = ["peso", "accion", "var_osc_pct", "var_rms_pct", "var_pico_pct", "nota"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h).font = FUENTE_NEGRITA

    if not conclusiones_path.exists():
        ws.cell(row=2, column=1, value="(fase2_sensibilidad.py no se ha corrido todavía)")
        return

    with open(conclusiones_path, "r", encoding="utf-8") as f:
        conclusiones = json.load(f)

    fila_out = 2
    for peso, c in conclusiones.items():
        fila_vals = [peso, c.get("accion"), c.get("var_osc_pct"), c.get("var_rms_pct"),
                     c.get("var_pico_pct"), c.get("nota")]
        for j, v in enumerate(fila_vals, start=1):
            ws.cell(row=fila_out, column=j, value=v)
        fila_out += 1


def main():
    seleccion_path = config.RESULTADOS_DIR / "fase4_seleccion.json"
    if not seleccion_path.exists():
        raise RuntimeError(f"No existe {seleccion_path}; correr fase4_seleccion.py primero.")
    with open(seleccion_path, "r", encoding="utf-8") as f:
        seleccion = json.load(f)

    wb = openpyxl.load_workbook(config.EXCEL_RESULTADOS)
    ws = wb["Pruebas realizadas"]

    print("=" * 60)
    print("  Fase 5: escribiendo Excel")
    print("=" * 60)

    for k, h in enumerate(config.HORIZONTES):
        h_str = str(h)
        if h_str not in seleccion:
            print(f"  Horizonte {h}: sin selección (saltado, correr Fase 4 primero)")
            continue
        print(f"  Horizonte {h}: escribiendo bloque {k} (fila {3+15*k})...")
        escribir_bloque(ws, k, h, seleccion[h_str])

    escribir_hoja_metricas_adicionales(wb, seleccion)
    escribir_hoja_resumen(wb, seleccion)
    escribir_hoja_sensibilidad(wb)

    wb.save(config.EXCEL_RESULTADOS)
    print(f"\nExcel guardado en {config.EXCEL_RESULTADOS}")
    print("Fase 5 (Excel) completa.")


if __name__ == "__main__":
    main()
